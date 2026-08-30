import os
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("GROQ_API_KEY not found in environment variables")

client = Groq(api_key=my_api_key)

model="openai/gpt-oss-20b"

# part 1
job_description="""
Senior Software Engineer — Backend Infrastructure

About the Role
We are looking for a Senior Software Engineer to join our Infrastructure team, 
building and scaling systems that serve millions of users worldwide. You will 
design distributed systems, drive technical strategy, and mentor engineers 
while working cross-functionally with Product, Design, and Data Science teams.

What You'll Do
- Design, build, and scale backend services handling high-throughput, 
  low-latency workloads across distributed systems
- Own the full lifecycle of features — from technical design and 
  implementation to deployment, monitoring, and iteration
- Partner with Product and Design to translate business requirements into 
  robust technical solutions
- Improve system reliability, performance, and observability through 
  better tooling, testing, and monitoring practices
- Participate in architecture reviews, design discussions, and code reviews 
  to maintain high engineering standards
- Mentor junior and mid-level engineers, and contribute to a strong 
  engineering culture
- Debug and resolve complex, production-level issues across the stack
- Contribute to technical roadmaps and long-term platform strategy

Minimum Qualifications
- Bachelor's degree in Computer Science, Engineering, or equivalent 
  practical experience
- 4+ years of experience in software development with a strong CS 
  fundamentals background (data structures, algorithms, system design)
- Proficiency in one or more of: Java, Python, Go, C++, or similar 
  backend languages
- Experience designing and building large-scale distributed systems
- Strong understanding of databases (SQL/NoSQL), caching, and 
  API design principles
- Experience with cloud infrastructure (AWS/GCP/Azure) and CI/CD pipelines

Preferred Qualifications
- Experience with microservices architecture and containerization 
  (Docker, Kubernetes)
- Familiarity with message queues/streaming systems (Kafka, RabbitMQ, etc.)
- Track record of leading technical projects end-to-end
- Prior experience at a high-scale consumer tech company
- Contributions to open-source projects or technical publications

What We Offer
- Competitive base salary + equity (RSUs) + annual performance bonus
- Comprehensive health, dental, and vision coverage
- Flexible work arrangements (hybrid/remote options)
- Learning & development budget for conferences, courses, certifications
- Relocation assistance (if applicable)
- 401(k)/retirement plan with company match

Location: Austin, TX | Employment Type: Full-Time
TechFlow Solutions is an Equal Opportunity Employer. We celebrate diversity and 
are committed to creating an inclusive environment for all employees.
"""

from pydantic import BaseModel

class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_qualifications: list[str]
    responsibilities: list[str]

jobd_schema=JobD.model_json_schema()

system_prompt=f"""
You are an expert HR Assistant. Your task is to read a job description 
and extract its details into a JSON object.

The JSON object must strictly follow this schema:

{jobd_schema}

Rules:
- Output ONLY the JSON object, nothing else (no markdown, no explanation).
- Fill each field based on what's mentioned in the job description.
- If some information is not present, use an empty list [] or null.
"""


user_prompt=f"""
Analyze the following job description and extract the information in the specified schema:

{job_description}
"""

system_message = {"role": "system", "content": system_prompt}
user_message = {"role": "user", "content": user_prompt}

messages = [system_message, user_message]

response_format={"type": "json_object"}

response = client.chat.completions.create(
    model=model,
    messages=messages,
    response_format=response_format
)

job_data = response.choices[0].message.content


import json

job_data_dict=json.loads(job_data)
job=JobD(**job_data_dict)


# part 2

class ResumeScore(BaseModel):
    score: float
    reason: dict

class Experience(BaseModel):
    company: str | None=None
    role: str | None=None
    duration: str | None=None
    description: str | None=None
    skills_used: list[str] = []

class Resume(BaseModel):
    name: str | None=None
    email: str | None=None
    phone: str | None=None

    total_experience_years: float | None=None

    skills: list[str]=[]
    experiences:list[Experience]=[]
    education:list[str]=[]
    projects:list[str]=[]
    certifications:list[str]=[]

resume_schema=Resume.model_json_schema()


# part 3

from pypdf import PdfReader
from docx import Document

def read_pdf(file_path: str) -> str:
    reader=PdfReader(file_path)
    text=""
    for page in reader.pages:
        if page.extract_text():
            text+=page.extract_text()+"\n"
    return text

def read_docx(file_path: str) -> str:
    doc=Document(file_path)
    text=""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text+=paragraph.text+"\n"
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text+=cell.text+"\n"
    return text

def read_resume(file_path: str) -> str:
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None


def parse_resume(resume_text: str) -> Resume:
    system_prompt=f"""
    You are an expert Resume Parser. Your task is to read the full text of a 
    resume and extract its details into a JSON object.

    The JSON object must strictly follow this schema:

    {resume_schema}

    Rules:
    - Output ONLY the JSON object, nothing else (no markdown, no explanation).
    - Fill each field based on what's mentioned in the resume text.
    - If some information is not present, use null (for single values) or an 
    empty list [] (for lists).
    - Each entry in "experiences" should correspond to one distinct job/role.
    """

    user_prompt=f"Parse the following resume:\n{resume_text}"

    system_message={
        "role":"system",
        "content":system_prompt
    }

    user_message={
        "role":"user",
        "content":user_prompt
    }

    response_format={"type":"json_object"}

    messages=[system_message, user_message]
    response=client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format
    )

    raw_response = response.choices[0].message.content
    parsed_resume = json.loads(raw_response)
    resume = Resume(**parsed_resume)
    return resume


def final_resume_score(job: JobDescription, resume: Resume) -> ResumeScore:
    resume_score_schema = ResumeScore.model_json_schema()
    system_prompt=f"""
    You are an expert Resume Scorer. Your task is to evaluate how well a candidate's 
    resume matches the given job description.

    The JSON object must strictly follow this schema:

    {resume_score_schema}

    Rules:
    - Output ONLY the JSON object, nothing else (no markdown, no explanation).
    - Calculate a score from 0 to 100 based on:
    - Relevance of skills, experience, and education to the job.
    - Strength of match (e.g., required vs preferred qualifications).
    - Overall alignment with the job requirements.
    - Provide a brief reason (1-2 sentences) explaining the score.
    """

    user_prompt=f"Job Description:\n{job.model_dump_json()}\n\nResume:\n{resume.model_dump_json()}"

    system_message={
        "role":"system",
        "content":system_prompt
    }

    user_message={
        "role":"user",
        "content":user_prompt
    }

    response_format={"type":"json_object"}

    messages=[system_message, user_message]
    response=client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format
    )

    raw_response = response.choices[0].message.content
    parsed_score = json.loads(raw_response)
    score = ResumeScore(**parsed_score)
    return score

# part 4

import time

resume_folder= Path("resumes")
all_resumes=[]

for file_path in resume_folder.iterdir():
    if file_path.is_file() and file_path.suffix.lower() in [".pdf", ".docx"]:
        print(f"Processing: {file_path.name}")
        resume_text = read_resume(file_path)
        parsed_resume=parse_resume(resume_text)
        time.sleep(5)
        result=final_resume_score(job, parsed_resume)
        time.sleep(5)
        all_resumes.append({
            "name":parsed_resume.name,
            "score":result.score,
            "reason":result.reason
        })

all_resumes.sort(key=lambda x: x["score"], reverse=True)


top_2_resumes = all_resumes[:2]
botton_2_resumes = all_resumes[-2:]

print("Top 2 Resumes:")
for resume in top_2_resumes:
    print(f"{resume['name']}: {resume['score']}%")
    print(f"Reason: {resume['reason']}")
    print()

print("Bottom 2 Resumes:")
for resume in botton_2_resumes:
    print(f"{resume['name']}: {resume['score']}%")
    print(f"Reason: {resume['reason']}")
    print()
